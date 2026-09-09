#!/usr/bin/env python3
"""Baseline UNIVERSAL de conciliación CFDI recibidos ↔ mpro — todos los
orígenes a la vez, chequeo agregado por CFDI (no por documento individual).

Motivación (2026-09-09): `baseline_conciliacion.py` concilia un origen a la
vez (hoy: COMPRA) exigiendo que UN documento cuadre exacto contra el CFDI.
Eso choca con un patrón real: un mismo CFDI puede estar repartido entre
VARIOS documentos de mpro — mismo origen (folio duplicado) o distinto
origen (COMPRA + COMPRA_INDIRECTO, GASTO_REGISTRO + CUENTA_X_PAGAR, o el
patrón de liquidación directa COMPRA cancelada + CHEQUE activo). Exigir 1
documento = 1 CFDI descarta esos casos aunque el dinero sí esté, solo que
repartido.

Este script no elige "el" documento: para cada CFDI SUMA el cargo de TODOS
los documentos con los que aparece etiquetado en `Comprobante_Digital`
(cualquier `Cd_Tabla`/origen), y compara esa suma contra el SUBTOTAL del
CFDI. Es un chequeo de un solo lado (cargo = reconocimiento de
inventario/gasto) — NO exige además que el abono cuadre; ver "SUPUESTOS"
en la hoja de salida.

Validado en vivo 2026-09-09, CFDI recibidos feb-2026: de 1,500 CFDI con
valor monetario real (subtotal > $1) encontrados en mpro, **1,305 (87.0%)**
ya cuadran en agregado — muy por encima del techo por-origen-individual
(81.8% COMPRA con el método de doble chequeo, ~37% Gasto_Registro con el
método viejo de `reconciliacion_por_origen.py`).

Uso:
    python3 baseline_universal.py --periodo 2026-02
    python3 baseline_universal.py --periodo 2026-02 --salida output/universal_feb2026.xlsx
"""
from __future__ import annotations

import argparse
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent / "src"))

import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from extract_sat import extract_sat_recibidos  # noqa: E402
from extract_origen import extract_origenes_por_uuids  # noqa: E402
from extract_poliza_por_origen import extract_poliza_por_origen, extract_poliza_cheque  # noqa: E402

TOL = 1.00
SUBTOTAL_MIN = 1.00  # bajo esto se considera CFDI sin valor monetario (TRASLADO/COMPROBANTE_PAGO)

# Orígenes cuyo Pd_Referencia no liga de forma confiable al folio del
# documento (ver poliza-explor/index.md) — se manejan aparte, no con el
# query genérico por Pd_Referencia=documento.
ORIGEN_MONTO_ESPECIAL = {"CHEQUE"}
# Orígenes que son "complementos" sin valor propio (Carta Porte, REP) — se
# excluyen de la búsqueda de cargo (no aportan, y consultarlos es ruido).
ORIGEN_SIN_VALOR = {"TRASLADO", "COMPROBANTE_PAGO"}

FONT = "Arial"
HEADER_FILL = PatternFill(start_color="2E4374", end_color="2E4374", fill_type="solid")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(name=FONT, bold=True, size=14, color="1A2233")
SUB_FONT = Font(name=FONT, size=10, italic=True, color="5B6472")
CELL_FONT = Font(name=FONT, size=10)
GOOD_FILL = PatternFill(start_color="E1F2EA", end_color="E1F2EA", fill_type="solid")
WARN_FILL = PatternFill(start_color="FBEEDA", end_color="FBEEDA", fill_type="solid")
THIN = Side(style="thin", color="D7DCE3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY_FMT = "$#,##0.00"
DATE_FMT = "yyyy-mm-dd"

COLS = {
    "uuid": ("UUID CFDI", 38),
    "fecha": ("Fecha", 12),
    "rfc_emisor": ("RFC Emisor", 14),
    "nombre_emisor": ("Proveedor", 32),
    "origenes": ("Orígen(es) en mpro", 26),
    "subtotal": ("Subtotal CFDI", 14),
    "iva": ("IVA CFDI", 12),
    "total": ("Total CFDI", 14),
    "cargo_agregado": ("Cargo agregado (todos los docs)", 18),
    "diferencia": ("Diferencia vs Subtotal", 16),
    "cuadra_via": ("Vía de cuadre", 20),
    "motivo_pendiente": ("Motivo pendiente", 20),
}


def calcular(periodo: str) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    print(f"[1/5] SAT recibidos {periodo}")
    sat = extract_sat_recibidos(periodo=periodo)
    print(f"     {len(sat)} CFDI recibidos")

    print("[2/5] Orígenes en Comprobante_Digital (TODOS, no solo un origen)")
    origenes = extract_origenes_por_uuids(sat["uuid"].tolist())
    origenes = origenes.dropna(subset=["documento"]).copy()
    origenes["documento_real"] = origenes["documento"].str.slice(0, 10)
    origenes["origen_up"] = origenes["origen"].str.upper()
    origenes = origenes.drop_duplicates(subset=["uuid", "origen_up", "documento_real"])
    print(f"     {origenes['uuid'].nunique()} CFDI con al menos 1 etiqueta en mpro"
          f" ({len(origenes)} etiquetas documento, algunos CFDI tienen varias)")

    print("[3/5] Cargo por documento, para cada origen presente (genérico, excluye cuentas de orden)")
    origenes_cargo = sorted(o for o in origenes["origen"].unique()
                             if o.upper() not in ORIGEN_MONTO_ESPECIAL | ORIGEN_SIN_VALOR)
    cargo_parts = []
    for origen in origenes_cargo:
        docs = origenes.loc[origenes["origen"] == origen, "documento_real"].dropna().unique().tolist()
        if not docs:
            continue
        pol = extract_poliza_por_origen(origen, docs)
        if pol.empty:
            continue
        pol = pol.rename(columns={"documento": "documento_real"})
        pol["origen"] = origen
        cargo_parts.append(pol[["origen", "documento_real", "cargo"]])
        print(f"     {origen}: {len(docs)} documentos -> {len(pol)} con cargo/abono en póliza")
    cargo_df = pd.concat(cargo_parts, ignore_index=True) if cargo_parts else pd.DataFrame(
        columns=["origen", "documento_real", "cargo"])

    merged = origenes.merge(cargo_df, on=["origen", "documento_real"], how="left")
    merged["cargo"] = merged["cargo"].fillna(0.0)
    total_cargo = merged.groupby("uuid")["cargo"].sum().rename("cargo_agregado")

    print("[4/5] CHEQUE: match por monto (Pd_Referencia no es confiable para este origen)")
    docs_cheque = origenes.loc[origenes["origen_up"] == "CHEQUE", "documento_real"].dropna().unique().tolist()
    if docs_cheque:
        cheque = extract_poliza_cheque(docs_cheque)
        cheque_map = origenes[origenes["origen_up"] == "CHEQUE"][["uuid", "documento_real"]].merge(
            cheque.rename(columns={"documento": "documento_real"})[["documento_real", "abono"]],
            on="documento_real", how="left")
        cheque_abono = cheque_map.groupby("uuid")["abono"].sum().rename("cheque_abono")
    else:
        cheque_abono = pd.Series(dtype=float, name="cheque_abono")

    print("[5/5] Armando resultado")
    origenes_str = origenes.groupby("uuid")["origen"].apply(lambda s: ", ".join(sorted(set(s)))).rename("origenes")

    resumen = sat.set_index("uuid").join(total_cargo, how="left").join(cheque_abono, how="left").join(
        origenes_str, how="left").reset_index()
    resumen["cargo_agregado"] = resumen["cargo_agregado"].fillna(0.0)
    resumen["cheque_abono"] = resumen["cheque_abono"].fillna(0.0)
    resumen["en_mpro"] = resumen["uuid"].isin(set(origenes["uuid"]))
    resumen["monetario"] = resumen["subtotal"].abs() > SUBTOTAL_MIN

    resumen["cuadra_cargo_subtotal"] = (resumen["cargo_agregado"] - resumen["subtotal"]).abs() <= TOL
    resumen["cuadra_pago_directo"] = (~resumen["cuadra_cargo_subtotal"]) & (
        (resumen["cheque_abono"] - resumen["total"]).abs() <= TOL)
    resumen["cuadra_agregado"] = resumen["cuadra_cargo_subtotal"] | resumen["cuadra_pago_directo"]

    def via(row):
        if row["cuadra_cargo_subtotal"]:
            return "cargo = subtotal"
        if row["cuadra_pago_directo"]:
            return "pago directo (Cheque = total)"
        return ""

    def motivo(row):
        if row["cuadra_agregado"]:
            return "OK"
        if not row["en_mpro"]:
            return "no encontrado en mpro"
        return "cargo agregado ≠ subtotal"

    resumen["cuadra_via"] = resumen.apply(via, axis=1)
    resumen["motivo_pendiente"] = resumen.apply(motivo, axis=1)
    resumen["diferencia"] = (resumen["cargo_agregado"] - resumen["subtotal"]).round(2)

    universo = resumen[resumen["monetario"] & resumen["en_mpro"]].copy()
    cols_base = ["uuid", "fecha", "rfc_emisor", "nombre_emisor", "origenes",
                 "subtotal", "iva", "total", "cargo_agregado", "diferencia"]

    conciliados = universo[universo["cuadra_agregado"]][cols_base + ["cuadra_via"]].copy()
    pendientes = universo[~universo["cuadra_agregado"]][cols_base + ["motivo_pendiente"]].copy()

    return conciliados, pendientes, resumen


def escribe_hoja(ws, df, cols, fill_status=None):
    headers = [COLS[c][0] for c in cols]
    ws.append(headers)
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=j)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(j)].width = COLS[c][1]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"

    money_cols = {"subtotal", "iva", "total", "cargo_agregado", "diferencia"}
    for _, row in df.iterrows():
        vals = []
        for c in cols:
            v = row.get(c)
            if pd.isna(v):
                v = "" if c not in money_cols else 0
            vals.append(v)
        ws.append(vals)
        r = ws.max_row
        for j, c in enumerate(cols, start=1):
            cell = ws.cell(row=r, column=j)
            cell.font = CELL_FONT
            cell.border = BORDER
            if c in money_cols:
                cell.number_format = MONEY_FMT
            if c == "fecha":
                cell.number_format = DATE_FMT
            if fill_status is not None:
                cell.fill = fill_status


def hoja_portada(ws, periodo, conciliados, pendientes, resumen):
    ws.column_dimensions["A"].width = 96
    ws["A1"] = "Conciliación CFDI ↔ mpro — Baseline UNIVERSAL (recibidos, todos los orígenes)"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Periodo: {periodo}   ·   Generado: {datetime.now().strftime('%Y-%m-%d')}"
    ws["A2"].font = SUB_FONT

    total_universo = len(conciliados) + len(pendientes)
    n_total_cfdi = len(resumen)
    n_en_mpro = int(resumen["en_mpro"].sum())
    n_sin_valor = int((resumen["en_mpro"] & ~resumen["monetario"]).sum())

    filas = [
        "",
        "MÉTODO",
        "A diferencia de los baselines por origen (que exigen que UN documento cuadre exacto contra el CFDI),",
        "este chequeo SUMA el cargo de TODOS los documentos de mpro con los que un CFDI aparece etiquetado en",
        "Comprobante_Digital, sin importar el origen (Cd_Tabla) — y compara esa suma contra el SUBTOTAL del CFDI",
        "(tolerancia $1.00). Así, un CFDI repartido entre COMPRA + COMPRA_INDIRECTO, o entre GASTO_REGISTRO +",
        "CUENTA_X_PAGAR, o duplicado bajo dos folios del mismo origen, puede cuadrar sin necesidad de elegir 'el'",
        "documento correcto — se suman todos.",
        "",
        "Caso especial CHEQUE: Pd_Referencia no liga de forma confiable al folio (72% de las veces trae la",
        "referencia externa del beneficiario, no el folio del cheque — ver poliza-explor). Para CFDI etiquetados",
        "bajo CHEQUE se usa match por MONTO (± $1 contra Ch_Importe) en vez de por referencia, y se compara contra",
        "el TOTAL del CFDI (no el subtotal) — cubre el patrón de 'liquidación directa' (proveedor paga de contado,",
        "la póliza de COMPRA se cancela y el pago se recaptura directo en Cheque).",
        "",
        "Se excluyen de la búsqueda de cargo TRASLADO y COMPROBANTE_PAGO (complementos SAT — Carta Porte y REP —",
        "sin valor propio, Cd_Monto siempre $0 en mpro) y, del universo comparado, cualquier CFDI con Subtotal ≤ $1.",
        "",
        "IMPORTANTE — esto es un chequeo de UN SOLO LADO (cargo = reconocimiento de inventario/gasto). NO exige",
        "que el abono (pago) también cuadre, a diferencia de baseline_conciliacion.py (COMPRA, doble chequeo).",
        "Un CFDI 'conciliado' aquí tiene su compra/gasto correctamente reconocido en la póliza, pero el lado del",
        "pago puede seguir sin verificar — ese es el trabajo pendiente, origen por origen.",
        "",
        "UNIVERSO",
        f"  • {n_total_cfdi} CFDI recibidos en el periodo.",
        f"  • {n_en_mpro} encontrados en mpro (al menos 1 etiqueta en Comprobante_Digital).",
        f"  • {n_sin_valor} de esos son complementos sin valor monetario (Subtotal ≤ $1 — típicamente TRASLADO/COMPROBANTE_PAGO) — se excluyen del cuadre.",
        f"  • {total_universo} CFDI con valor monetario real, encontrados en mpro — este es el universo comparado abajo.",
        "",
        "RESULTADO",
        f"  • Conciliados (cargo agregado cuadra, vía cargo=subtotal o pago directo): {len(conciliados)} CFDI  ({len(conciliados)/total_universo*100:.1f}%)" if total_universo else "  • Sin datos",
        f"  • Pendientes: {len(pendientes)} CFDI  ({len(pendientes)/total_universo*100:.1f}%)" if total_universo else "",
        "",
        "Por comparación: el baseline por origen (solo COMPRA, doble chequeo cargo+abono) da 81.8% sobre 713 CFDI;",
        "aquí, con el chequeo agregado de un solo lado mas COMPRA solo llega a 94.3% (701 CFDI) y Gasto_Registro a",
        "82.8% (699 CFDI) — muy por encima del ~37% que daba el método viejo (reconciliacion_por_origen.py) para",
        "ese origen, porque ahora la exclusión de 'cuentas de orden' usa Poliza_Configuracion (robusta) en vez del",
        "filtro de texto anterior.",
        "",
        "HOJAS",
        "  • Conciliados: CFDI + orígenes donde aparece + cargo agregado + vía de cuadre.",
        "  • Pendientes: CFDI encontrados en mpro pero sin cuadrar en agregado, con el motivo.",
    ]
    for i, texto in enumerate(filas, start=3):
        cell = ws.cell(row=i, column=1, value=texto)
        cell.font = Font(name=FONT, size=10, bold=(texto in ("MÉTODO", "UNIVERSO", "RESULTADO", "HOJAS")))
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def run(periodo: str, salida: str):
    conciliados, pendientes, resumen = calcular(periodo)
    total = len(conciliados) + len(pendientes)
    print(f"\nUniverso monetario en mpro: {total}   Conciliados: {len(conciliados)} "
          f"({len(conciliados)/total*100:.1f}%)   Pendientes: {len(pendientes)}")

    Path(salida).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Resumen"
    hoja_portada(ws0, periodo, conciliados, pendientes, resumen)

    cols1 = ["uuid", "fecha", "rfc_emisor", "nombre_emisor", "origenes", "subtotal", "iva", "total",
             "cargo_agregado", "diferencia", "cuadra_via"]
    ws1 = wb.create_sheet("Conciliados")
    escribe_hoja(ws1, conciliados, cols1, fill_status=GOOD_FILL)

    cols2 = ["uuid", "fecha", "rfc_emisor", "nombre_emisor", "origenes", "subtotal", "iva", "total",
             "cargo_agregado", "diferencia", "motivo_pendiente"]
    ws2 = wb.create_sheet("Pendientes")
    escribe_hoja(ws2, pendientes, cols2, fill_status=WARN_FILL)

    wb.save(salida)
    print(f"\nReporte: {salida}")
    return conciliados, pendientes, resumen


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--periodo", required=True)
    ap.add_argument("--salida", default=None)
    args = ap.parse_args()
    salida = args.salida or f"output/baseline_universal_{args.periodo}.xlsx"
    run(args.periodo, salida)
