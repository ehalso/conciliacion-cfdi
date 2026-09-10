"""Genera el reporte .xlsx de la conciliación: Resumen + Detalle."""
from __future__ import annotations

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ESTATUS_COLORS = {
    "OK": "C6EFCE",
    "CONCILIADO": "C6EFCE",
    "DIFERENCIA_IMPORTE": "FFC7CE",
    "SOLO_SAT": "FFEB9C",
    "SOLO_MPRO": "FFEB9C",
    "STUB_GASTO_REGISTRO_SIN_MONTO": "FFEB9C",
    "SIN_MAPEO_MPRO": "FFC7CE",
}


def write_report(detalle: pd.DataFrame, resumen_df: pd.DataFrame, salida: str) -> None:
    with pd.ExcelWriter(salida, engine="openpyxl") as writer:
        resumen_df.to_excel(writer, sheet_name="Resumen", index=False)
        detalle.to_excel(writer, sheet_name="Detalle", index=False)

        for field in (
            "SOLO_SAT",
            "SOLO_MPRO",
            "DIFERENCIA_IMPORTE",
            "STUB_GASTO_REGISTRO_SIN_MONTO",
            "SIN_MAPEO_MPRO",
        ):
            subset = detalle[detalle["estatus"] == field]
            if not subset.empty:
                sheet_name = field[:31]
                subset.to_excel(writer, sheet_name=sheet_name, index=False)

        wb = writer.book
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for col_idx, col_cells in enumerate(ws.columns, start=1):
                max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)
            for cell in ws[1]:
                cell.font = Font(bold=True)

            if "estatus" in [c.value for c in ws[1]]:
                estatus_col = [c.value for c in ws[1]].index("estatus") + 1
                for row in ws.iter_rows(min_row=2, min_col=estatus_col, max_col=estatus_col):
                    for cell in row:
                        color = ESTATUS_COLORS.get(cell.value)
                        if color:
                            cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
