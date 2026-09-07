#!/usr/bin/env python3
"""Conciliación a nivel póliza/cuenta contable: para cada CFDI recibido,
trazabilidad hasta la póliza y cuenta contable real que lo registró en
mpro, vía Poliza_Detalle_Comprobante (ver src/extract_poliza.py para el
razonamiento y las decisiones de alcance).

Uso:
    python3 poliza_reconciliation.py --periodo 2026-01 --salida output/poliza_2026-01.xlsx
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent / "src"))

from extract_sat import extract_sat_recibidos  # noqa: E402
from extract_poliza import extract_poliza_por_uuids  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

TOLERANCIA = 1.00


def clasificar(row) -> str:
    if not row["con_poliza"]:
        return "SIN_HUELLA_CONTABLE"
    cfdi_total = row["total"]
    if abs(row["suma_cargo"] - cfdi_total) <= TOLERANCIA or abs(row["suma_abono"] - cfdi_total) <= TOLERANCIA:
        return "CARGO_O_ABONO_COINCIDE_CON_TOTAL"
    if row["suma_cargo"] > 0 or row["suma_abono"] > 0:
        return "CON_HUELLA_SIN_COINCIDIR_TOTAL"
    return "SIN_HUELLA_CONTABLE"


def run(periodo: str, salida: str):
    Path(salida).parent.mkdir(parents=True, exist_ok=True)

    print(f"[1/3] Extrayendo SAT recibidos periodo={periodo} ...")
    sat_df = extract_sat_recibidos(periodo=periodo)
    print(f"      {len(sat_df)} filas")
    if sat_df.empty:
        print("Sin datos, nada que hacer.")
        return

    print(f"[2/3] Trazando {sat_df['uuid'].nunique()} UUIDs hasta póliza/cuenta contable (Poliza_Detalle_Comprobante, .207) ...")
    pol_df = extract_poliza_por_uuids(sat_df["uuid"].tolist())
    print(f"      {len(pol_df)} UUIDs con huella contable activa de {sat_df['uuid'].nunique()} totales")

    merged = sat_df.merge(pol_df, on="uuid", how="left")
    merged["con_poliza"] = merged["con_poliza"].fillna(False)
    for c in ("suma_cargo", "suma_abono", "n_polizas_activas", "n_cuentas_distintas"):
        merged[c] = merged[c].fillna(0)
    merged["estatus_contable"] = merged.apply(clasificar, axis=1)

    print("[3/3] Resumen:")
    resumen = merged.groupby("estatus_contable").agg(
        conteo=("uuid", "count"),
        monto_total_cfdi=("total", "sum"),
    ).reset_index()
    print(resumen.to_string(index=False))

    cols = [
        "uuid", "estatus_contable", "tipo_comprobante", "rfc_emisor",
        "subtotal", "iva", "total",
        "con_poliza", "n_polizas_activas", "n_cuentas_distintas",
        "suma_cargo", "suma_abono", "cuentas_cargo", "cuentas_abono", "periodo",
    ]
    detalle = merged[cols]

    with pd.ExcelWriter(salida, engine="openpyxl") as writer:
        resumen.to_excel(writer, sheet_name="Resumen", index=False)
        detalle.to_excel(writer, sheet_name="Detalle", index=False)
        for est in detalle["estatus_contable"].unique():
            sub = detalle[detalle["estatus_contable"] == est]
            if not sub.empty:
                sub.to_excel(writer, sheet_name=est[:31], index=False)

        wb = writer.book
        colors = {
            "CARGO_O_ABONO_COINCIDE_CON_TOTAL": "C6EFCE",
            "CON_HUELLA_SIN_COINCIDIR_TOTAL": "FFEB9C",
            "SIN_HUELLA_CONTABLE": "FFC7CE",
        }
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for col_idx, col_cells in enumerate(ws.columns, start=1):
                max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)
            for cell in ws[1]:
                cell.font = Font(bold=True)
            headers = [c.value for c in ws[1]]
            if "estatus_contable" in headers:
                idx = headers.index("estatus_contable") + 1
                for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                    for cell in row:
                        color = colors.get(cell.value)
                        if color:
                            cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")

    print(f"Reporte: {salida}")
    return merged, resumen


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--periodo", required=True)
    ap.add_argument("--salida", default=None)
    args = ap.parse_args()
    salida = args.salida or f"output/poliza_{args.periodo}.xlsx"
    run(args.periodo, salida)
