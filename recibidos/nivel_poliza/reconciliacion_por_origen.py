#!/usr/bin/env python3
"""Reconciliación de CFDI recibidos contra mpro, distinguida por origen de
documento (Comprobante_Digital.Cd_Tabla) — para saber exactamente qué
módulos de mpro hay que reconciliar y cuánto de cada uno ya cuadra.

Uso: python3 reconciliacion_por_origen.py --periodo 2026-02
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent.parent / "src"))

import pandas as pd  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from extract_sat import extract_sat_recibidos  # noqa: E402
from extract_origen import extract_origenes_por_uuids  # noqa: E402
from extract_poliza_por_origen import extract_poliza_por_origen, extract_poliza_cheque  # noqa: E402

TOLERANCIA = 1.00

# Orígenes con volumen/monto relevante y mapeo confirmado a Poliza_Control.
# TRASLADO se excluye a propósito: 100% CFDI tipo T, monto $0 (Carta Porte
# autoemitida, sin efecto fiscal) — confirmado con datos, no aporta a la
# reconciliación de importes.
ORIGENES_A_RECONCILIAR = [
    "COMPRA", "GASTO_REGISTRO", "CUENTA_X_PAGAR", "CHEQUE",
    "NOTA_CREDITO_PROVEEDOR", "COMPRA_INDIRECTO",
]


def run(periodo: str, salida: str):
    Path(salida).parent.mkdir(parents=True, exist_ok=True)

    print(f"[1/4] Extrayendo SAT recibidos {periodo} ...")
    sat_df = extract_sat_recibidos(periodo=periodo)
    print(f"      {len(sat_df)} CFDI")

    print("[2/4] Trazando origen (Comprobante_Digital.Cd_Tabla/Cd_Documento) ...")
    origenes = extract_origenes_por_uuids(sat_df["uuid"].tolist())
    print(f"      {len(origenes)} filas origen para {origenes['uuid'].nunique()} CFDI matcheados")

    sat_amounts = sat_df.set_index("uuid")[["subtotal", "iva", "total", "tipo_comprobante"]]

    all_detail = []
    resumen_rows = []
    doc_level_rows = []

    for origen in ORIGENES_A_RECONCILIAR:
        sub = origenes[origenes["origen"] == origen].drop_duplicates(subset=["uuid", "documento"]).copy()
        if sub.empty:
            continue
        # Gotcha confirmado en vivo 2026-09-07 (mismo patrón ya documentado
        # para GASTO_REGISTRO en layout-gastos): Comprobante_Digital.Cd_Documento
        # trae el folio real ('XX-NNNNNNN', 10 caracteres) + un sufijo de
        # 4-8 dígitos (línea/sub-documento) que Poliza_Control.Pc_Documento
        # NO lleva. Se trunca a los primeros 10 caracteres antes de buscar
        # la póliza. Verificado con muestra real de los 6 orígenes.
        sub["documento_real"] = sub["documento"].str.slice(0, 10)
        sub = sub.join(sat_amounts, on="uuid")
        docs = sub["documento_real"].dropna().unique().tolist()
        print(f"[3/4] {origen}: {sub['uuid'].nunique()} CFDI, {len(docs)} documentos (folio real) ...")
        if origen == "CHEQUE":
            pol = extract_poliza_cheque(docs).rename(columns={"documento": "documento_real"})
        else:
            pol = extract_poliza_por_origen(origen, docs).rename(columns={"documento": "documento_real"})

        # Comparación a nivel DOCUMENTO, no CFDI: un mismo folio de mpro
        # (ej. un cheque, una compra) puede consolidar varios CFDI — el
        # cargo/abono posteado cubre la suma del folio, no un CFDI aislado.
        # Confirmado con datos: el folio '01-0079502' de Cheque cubre 3
        # UUID distintos. Se agrega primero, se compara después.
        doc_group = sub.groupby("documento_real").agg(
            n_cfdi=("uuid", "nunique"),
            monto_cfdi=("total", "sum"),
            monto_subtotal=("subtotal", "sum"),
        ).reset_index()
        doc_group = doc_group.merge(pol, on="documento_real", how="left")
        doc_group["cargo"] = pd.to_numeric(doc_group["cargo"], errors="coerce").fillna(0.0)
        doc_group["abono"] = pd.to_numeric(doc_group["abono"], errors="coerce").fillna(0.0)
        doc_group["n_polizas"] = pd.to_numeric(doc_group["n_polizas"], errors="coerce").fillna(0).astype(int)

        # Hallazgo confirmado en vivo 2026-09-07: el Cargo aislado por
        # Pd_Referencia NO siempre cuadra contra el TOTAL del CFDI — para
        # varios orígenes (COMPRA sobre todo, también buena parte de
        # GASTO_REGISTRO y CUENTA_X_PAGAR) el Cargo se postea al SUBTOTAL
        # (cuenta de inventario/gasto, sin IVA) porque el IVA se registra en
        # una cuenta/línea aparte que no lleva el mismo Pd_Referencia. Se
        # acepta como "cuadra" si el cargo o abono coincide con el TOTAL *o*
        # con el SUBTOTAL — se guarda contra cuál de los dos para que quede
        # trazable en el detalle, no oculto.
        #
        # CHEQUE es un caso aparte: el origen liga mayoritariamente CFDI de
        # tipo P (Recibo Electrónico de Pago / REP) cuyo Subtotal y Total
        # son $0.00 por diseño del SAT (el monto real vive en el nodo
        # Complemento de Pagos, que este pipeline no parsea) — confirmado en
        # vivo: 654 de 665 filas de Cheque en febrero son tipo P con
        # subtotal=iva=total=0. Por eso Cheque NO se compara contra
        # subtotal/total del CFDI: extract_poliza_cheque() ya solo devuelve
        # una póliza cuando `ABS(Pd_Importe - Ch_Importe) <= 1`, es decir el
        # match de monto va implícito en la extracción misma. Aquí, tener
        # n_polizas>0 en Cheque YA significa que cuadra.
        def clasif(r):
            if r["n_polizas"] == 0:
                return "SIN_POLIZA", ""
            if origen == "CHEQUE":
                return "CUADRA", "MONTO_CHEQUE_VS_ABONO"
            if abs(r["cargo"] - r["monto_cfdi"]) <= TOLERANCIA or abs(r["abono"] - r["monto_cfdi"]) <= TOLERANCIA:
                return "CUADRA", "TOTAL"
            if abs(r["cargo"] - r["monto_subtotal"]) <= TOLERANCIA or abs(r["abono"] - r["monto_subtotal"]) <= TOLERANCIA:
                return "CUADRA", "SUBTOTAL"
            return "CON_POLIZA_SIN_CUADRAR", ""

        doc_group[["estatus", "base_cuadre"]] = doc_group.apply(lambda r: pd.Series(clasif(r)), axis=1)
        doc_group["origen"] = origen

        merged = sub.merge(doc_group[["documento_real", "cargo", "abono", "n_polizas", "estatus", "base_cuadre", "n_cfdi", "monto_cfdi", "monto_subtotal"]], on="documento_real", how="left")
        merged["origen"] = origen
        all_detail.append(merged)
        doc_level_rows.append(doc_group)

        g = merged.groupby("estatus").agg(n_cfdi=("uuid", "nunique"), monto_total=("total", "sum")).reset_index()
        for _, r in g.iterrows():
            resumen_rows.append({"origen": origen, "estatus": r["estatus"], "n_cfdi": r["n_cfdi"], "monto_total": r["monto_total"]})

    detalle = pd.concat(all_detail, ignore_index=True)
    resumen = pd.DataFrame(resumen_rows)
    doc_level = pd.concat(doc_level_rows, ignore_index=True)

    print("[4/4] Resumen por origen (a nivel CFDI, documentos consolidados):")
    piv = resumen.pivot_table(index="origen", columns="estatus", values="n_cfdi", fill_value=0, aggfunc="sum")
    print(piv.to_string())
    print("\nResumen por origen (a nivel DOCUMENTO mpro):")
    piv_doc = doc_level.pivot_table(index="origen", columns="estatus", values="documento_real", fill_value=0, aggfunc="count")
    print(piv_doc.to_string())

    # Cuadre AGREGADO por origen: la base son los documentos de mpro que YA
    # tienen XML/CFDI adjunto (Comprobante_Digital) — los que no tienen XML
    # se quedan fuera por ahora, según lo pedido. Aquí NO se exige que cada
    # documento cuadre individualmente contra su CFDI; se suma el cargo/abono
    # contabilizado de TODA la base y se compara contra la suma de
    # subtotal/total de los CFDI de esa misma base — para ver si el bloque
    # completo de un origen cuadra en conjunto aunque el detalle documento a
    # documento no sea perfecto (splits por centro de costo, notas partidas
    # en 2 folios, etc. se cancelan al sumar).
    cuadre_agregado = doc_level.groupby("origen").agg(
        n_docs=("documento_real", "count"),
        n_cfdi=("n_cfdi", "sum"),
        suma_subtotal_cfdi=("monto_subtotal", "sum"),
        suma_total_cfdi=("monto_cfdi", "sum"),
        suma_cargo=("cargo", "sum"),
        suma_abono=("abono", "sum"),
    ).reset_index()
    cuadre_agregado["pct_cargo_vs_subtotal"] = (
        cuadre_agregado["suma_cargo"] / cuadre_agregado["suma_subtotal_cfdi"] * 100
    ).round(1)
    cuadre_agregado["pct_cargo_vs_total"] = (
        cuadre_agregado["suma_cargo"] / cuadre_agregado["suma_total_cfdi"] * 100
    ).round(1)
    cuadre_agregado["pct_abono_vs_total"] = (
        cuadre_agregado["suma_abono"] / cuadre_agregado["suma_total_cfdi"] * 100
    ).round(1)
    cuadre_agregado["nota"] = ""
    cuadre_agregado.loc[cuadre_agregado["origen"] == "CHEQUE", "nota"] = (
        "Subtotal/Total del CFDI = $0 (tipo P, Recibo de Pago) — no comparable "
        "contra el CFDI por este medio; ver reporte, el cuadre real de Cheque "
        "es por monto (Ch_Importe vs Pd_Importe), no por CFDI."
    )
    print("\nCuadre AGREGADO por origen (suma total contabilizada vs suma total CFDI):")
    print(cuadre_agregado.to_string(index=False))

    cols = ["uuid", "origen", "documento", "documento_real", "estatus", "base_cuadre", "tipo_comprobante", "subtotal", "iva", "total",
            "n_cfdi", "monto_cfdi", "monto_subtotal", "cargo", "abono", "n_polizas"]
    detalle = detalle[cols]

    with pd.ExcelWriter(salida, engine="openpyxl") as writer:
        cuadre_agregado.to_excel(writer, sheet_name="Cuadre_Agregado", index=False)
        resumen.to_excel(writer, sheet_name="Resumen_por_origen_CFDI", index=False)
        doc_level.to_excel(writer, sheet_name="Resumen_por_documento", index=False)
        detalle.to_excel(writer, sheet_name="Detalle", index=False)
        for origen in ORIGENES_A_RECONCILIAR:
            sub = detalle[detalle["origen"] == origen]
            if not sub.empty:
                sub.to_excel(writer, sheet_name=origen[:31], index=False)

        wb = writer.book
        colors = {"CUADRA": "C6EFCE", "CON_POLIZA_SIN_CUADRAR": "FFEB9C", "SIN_POLIZA": "FFC7CE"}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for col_idx, col_cells in enumerate(ws.columns, start=1):
                max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)
            for cell in ws[1]:
                cell.font = Font(bold=True)
            headers = [c.value for c in ws[1]]
            if "estatus" in headers:
                idx = headers.index("estatus") + 1
                for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                    for cell in row:
                        color = colors.get(cell.value)
                        if color:
                            cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")

    print(f"\nReporte: {salida}")
    return detalle, resumen, cuadre_agregado


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--periodo", required=True)
    ap.add_argument("--salida", default=None)
    args = ap.parse_args()
    salida = args.salida or f"output/reconciliacion_origen_{args.periodo}.xlsx"
    run(args.periodo, salida)
